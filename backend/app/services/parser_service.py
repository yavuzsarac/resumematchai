from io import BytesIO

import fitz
from docx import Document

from app.services.text_cleaner import clean_text


class ResumeParsingError(ValueError):
    pass


def _parse_pdf(content: bytes) -> str:
    try:
        with fitz.open(stream=content, filetype="pdf") as document:
            return "\n".join(page.get_text("text") for page in document)
    except Exception as exc:
        raise ResumeParsingError(
            "The PDF could not be read. It may be damaged or password protected."
        ) from exc


def _parse_docx(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_cells = [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
        return "\n".join(paragraphs + table_cells)
    except Exception as exc:
        raise ResumeParsingError("The DOCX file could not be read.") from exc


def _parse_txt(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ResumeParsingError("The TXT file encoding is not supported.")


def parse_resume(content: bytes, extension: str) -> str:
    parsers = {
        ".pdf": _parse_pdf,
        ".docx": _parse_docx,
        ".txt": _parse_txt,
    }
    parser = parsers.get(extension.lower())
    if parser is None:
        raise ResumeParsingError("Unsupported resume file type.")

    cleaned = clean_text(parser(content))
    if len(cleaned) < 20:
        raise ResumeParsingError(
            "No readable resume text was found. Scanned PDFs require OCR and are not supported yet."
        )
    return cleaned

